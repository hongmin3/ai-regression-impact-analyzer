from datetime import datetime, timezone

from app.core.storage import Storage
from app.modules.impact_analyzer.ai_client import ImpactAnalysisAIClient
from app.modules.impact_analyzer.schemas import AnalysisResult, ChangeAnalysis, ChangeItem, DraftTestCase, Impact, ImpactDecision, SpecificationChunk, TestCase
from openpyxl import load_workbook

from app.modules.impact_analyzer.html_report import create_html_report, create_xlsx_export
from app.modules.impact_analyzer.tc_draft import create_tc_draft_markdown


def test_mock_gemini_structured_output(tmp_path):
    response = {
        "decisions": [{"tc_id": "TC-1", "impact": "HIGH", "confidence": .91, "direct_impact": True, "regression_needed": True, "reason": "저장 변경", "relevant_specifications": ["spec-p1-0"], "verification_points": ["재실행"], "recommended": True}],
        "draft_test_cases": [{"changed_feature": "신규 저장 옵션 추가", "evidence_chunk_ids": ["spec-p1-0"]}],
        "token_usage": {"prompt_tokens": 100, "candidate_tokens": 20, "total_tokens": 120},
    }
    client = ImpactAnalysisAIClient(Storage(tmp_path / "test.db"), responder=lambda _: response)
    values = client.analyze(ChangeAnalysis(changed_features=["설정 저장"]), [TestCase(tc_id="TC-1")], [SpecificationChunk(chunk_id="spec-p1-0", document_id="spec", page=1, text="설정 저장")])
    assert values[0].impact is Impact.HIGH
    assert client.request_count == 1
    assert client.token_usage == {"prompt_tokens": 100, "candidate_tokens": 20, "total_tokens": 120}
    assert client.draft_test_cases[0].changed_feature == "신규 저장 옵션 추가"


def test_html_report_hides_internal_fields_and_manual_review_section():
    result = AnalysisResult(
        analysis_id="report-fields-test", created_at=datetime.now(timezone.utc), change_file="change.pdf",
        specification_file="spec.pdf", testcase_file="tc.xlsx",
        change=ChangeAnalysis(change_items=[ChangeItem(feature="지문 인증 옵션 추가", impact_area="로그인")]),
        total_tc=2, candidate_tc=2,
        decisions=[
            ImpactDecision(tc_id="TC-1", impact="HIGH", confidence=1.0, reason="직접 영향", recommended=True, specification_reference="VXvue 사양서1 · p.10"),
            ImpactDecision(tc_id="TC-2", impact="LOW", confidence=.4, reason="정보 부족", recommended=True, manual_review_required=True),
        ],
    )
    text = create_html_report(result).read_text(encoding="utf-8")
    assert "Evidence Level" not in text
    assert "Revision Mark" not in text
    assert "Manual Review Required" not in text
    assert "Rev.1.7" not in text
    assert "확인 필요" in text
    assert "VXvue 사양서1 · p.10" in text
    assert "지문 인증 옵션 추가" in text


def test_tc_draft_markdown_written_with_confirm_needed_placeholders():
    result = AnalysisResult(
        analysis_id="draft-test", created_at=datetime.now(timezone.utc), change_file="change.pdf",
        specification_file="spec.pdf", testcase_file="tc.xlsx", change=ChangeAnalysis(),
        total_tc=1, candidate_tc=1, decisions=[],
        draft_test_cases=[DraftTestCase(changed_feature="신규 기능 추가", evidence_chunk_ids=["spec-p1-0"])],
    )
    path = create_tc_draft_markdown(result)
    assert path is not None and path.exists()
    content = path.read_text(encoding="utf-8")
    assert "신규 기능 추가" in content
    assert "확인 필요" in content


def test_tc_draft_markdown_skipped_when_no_drafts():
    result = AnalysisResult(
        analysis_id="no-draft-test", created_at=datetime.now(timezone.utc), change_file="change.pdf",
        specification_file="spec.pdf", testcase_file="tc.xlsx", change=ChangeAnalysis(),
        total_tc=1, candidate_tc=1, decisions=[],
    )
    assert create_tc_draft_markdown(result) is None


def test_report_generation():
    result = AnalysisResult(analysis_id="unit-test", created_at=datetime.now(timezone.utc), change_file="change.pdf", specification_file="spec.pdf", testcase_file="tc.xlsx", change=ChangeAnalysis(changed_features=["설정"]), total_tc=1, candidate_tc=1, decisions=[ImpactDecision(tc_id="TC-1", impact="HIGH", confidence=.9, reason="변경 영향", recommended=True)])
    assert create_html_report(result).exists()
    xlsx = create_xlsx_export(result)
    assert xlsx.exists()
    sheet = load_workbook(xlsx, read_only=True).active
    assert sheet["A1"].value == "TC ID"
    assert sheet["A2"].value == "TC-1"


def test_ai_client_keeps_exact_prompt_and_response_for_audit(tmp_path):
    response = {"decisions": [], "draft_test_cases": [], "change_items": [], "token_usage": {"total_tokens": 7}}
    client = ImpactAnalysisAIClient(Storage(tmp_path / "audit.db"), responder=lambda _: response)

    client.analyze(ChangeAnalysis(user_notes="감사 기록"), [], [])
    audit = client.audit_snapshot

    assert '"user_notes": "감사 기록"' in audit["user_prompt"]
    assert audit["system_instruction"]
    assert audit["response"] == response
    assert audit["cache_hit"] is False
